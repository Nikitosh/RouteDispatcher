import datetime

import openpyxl
import pytest

from dispatch.files.readers import InputError, parse_minutes, parse_modes, parse_skills, read_brigades, read_orders
from dispatch.files.tables import read_table

ORDERS_CSV = """Заявка;Тип заявки BK;Тип заявки HD;Начало;Окончание;Район;Адрес;Гигабитное подключение
101;Подключение;Конвергенция абонента;17.08.2026 10:00;17.08.2026 12:00;Лефортово;Город Москва, ул.Юности, д. 1;Нет
102;Глобальная проблема;Авария;17.08.2026 0:01;17.08.2026 23:59;Кузьминки;Город Москва, ул.Юности, д. 2;Нет
Адрес офиса;г. Москва, ул Юных Ленинцев, д 83с 4
"""


def write(tmp_path, name, text, encoding='cp1251'):
    p = tmp_path / name
    p.write_bytes(text.encode(encoding))
    return str(p)


def test_orders_from_organizers_csv(tmp_path):
    f = read_orders(write(tmp_path, 'o.csv', ORDERS_CSV))
    assert f.office == 'г. Москва, ул Юных Ленинцев, д 83с 4'
    assert f.date == '17.08.2026'
    a, b = f.orders
    assert (a.type, a.window, a.service, a.priority, a.skill) == ('Подключение', (0, 120), 70, 2, 1)
    assert b.type == 'Авария' and b.window == (-599, 839) and b.earliest == 0
    assert a.need is None


def test_apartment_and_gigabit(tmp_path):
    text = ORDERS_CSV.replace('ул.Юности, д. 1;Нет', 'ул.Юности, д. 1, кв. 71;Да')
    a, b = read_orders(write(tmp_path, 'o.csv', text)).orders
    assert (a.apt, a.gigabit, a.house) == (71, True, 'Город Москва, ул.Юности, д. 1')
    assert (b.apt, b.gigabit) == (None, False)


def test_utf8_and_transport_column(tmp_path):
    text = ORDERS_CSV.replace('Гигабитное подключение', 'Гигабитное подключение;Транспорт').replace(';Нет\n101', ';Нет;авто\n101')
    lines = text.splitlines()
    lines[1] += ';авто или велосипед'
    lines[2] += ';'
    f = read_orders(write(tmp_path, 'o.csv', '\n'.join(lines), 'utf-8'))
    assert f.orders[0].need == frozenset({'car', 'bike'})
    assert f.orders[1].need is None


def test_orders_errors_point_to_line(tmp_path):
    bad = ORDERS_CSV.replace('Глобальная проблема', 'Что-то новое')
    with pytest.raises(InputError, match='строка 3: неизвестный тип заявки'):
        read_orders(write(tmp_path, 'o.csv', bad))
    with pytest.raises(InputError, match='Адрес офиса'):
        read_orders(write(tmp_path, 'o.csv', ORDERS_CSV.rsplit('Адрес офиса', 1)[0]))


def test_brigades(tmp_path):
    text = 'Бригада;Транспорт;Навыки;Старт\nБ1;Авто;Ремонт, Подключение и дозаказ;\nБ2;пешком;аварии;55.1, 37.2\n'
    b1, b2 = read_brigades(write(tmp_path, 'b.csv', text))
    assert (b1.mode, b1.skills, b1.start_address) == ('car', [0, 1], '')
    assert (b2.mode, b2.skills, b2.start_address) == ('foot', [2], '55.1, 37.2')
    with pytest.raises(InputError, match='строка 2: неизвестный транспорт «самокат»'):
        read_brigades(write(tmp_path, 'b.csv', 'Бригада;Транспорт;Навыки\nБ1;самокат;ремонт\n'))


def test_brigade_schedules(tmp_path):
    text = ('Бригада;Транспорт;Навыки;График;Смена\nБ1;Авто;ремонт;2/2;\nБ2;Авто;ремонт;5/2;\nБ3;Авто;ремонт;5/2;\n'
            'Б4;Авто;ремонт;5/2;12-21\nБ5;Авто;ремонт;;\n')
    path = write(tmp_path, 'b.csv', text)
    assert [b.shift_text for b in read_brigades(path)] == ['10:00–22:00', '10:00–19:00', '13:00–22:00', '12:00–21:00', '10:00–22:00']
    assert read_brigades(path, '5/2')[-1].schedule == '5/2'     # пустой график — из параметра
    with pytest.raises(InputError, match='неизвестный график «3/3»'):
        read_brigades(write(tmp_path, 'b.csv', 'Бригада;Транспорт;Навыки;График\nБ1;Авто;ремонт;3/3\n'))
    with pytest.raises(InputError):
        read_brigades(write(tmp_path, 'b.csv', 'Бригада;Транспорт;Навыки;Смена\nБ1;Авто;ремонт;8-17\n'))


def test_brigade_level_and_speed(tmp_path):
    text = ('Бригада;Транспорт;Навыки;Уровень;Темп\nБ1;Авто;ремонт;Стажёр;\nБ2;Авто;ремонт;senior;\nБ3;Авто;ремонт;;\n'
            'Б4;Авто;ремонт;junior;0,95\nБ5;Авто;ремонт;Middle;×1.1\nБ6;Авто;ремонт;старший инженер;110%\n')
    bs = read_brigades(write(tmp_path, 'b.csv', text, 'utf-8'))
    assert [b.level for b in bs] == ['стажёр', 'старший', '', 'стажёр', 'инженер', 'старший']
    assert [b.speed for b in bs] == [1.15, 0.9, 1.0, 0.95, 1.1, pytest.approx(1.1)]    # «Темп» важнее уровня
    assert all(b.speed == 1.0 and b.level == '' for b in read_brigades(write(tmp_path, 'b.csv', 'Бригада;Транспорт;Навыки\nБ1;Авто;ремонт\n')))
    with pytest.raises(InputError, match='строка 2: неизвестный уровень «гуру»'):
        read_brigades(write(tmp_path, 'b.csv', 'Бригада;Транспорт;Навыки;Уровень\nБ1;Авто;ремонт;гуру\n'))
    with pytest.raises(InputError, match='темп «5» вне пределов'):
        read_brigades(write(tmp_path, 'b.csv', 'Бригада;Транспорт;Навыки;Темп\nБ1;Авто;ремонт;5\n'))
    with pytest.raises(InputError, match='не понял темп'):
        read_brigades(write(tmp_path, 'b.csv', 'Бригада;Транспорт;Навыки;Темп\nБ1;Авто;ремонт;быстро\n'))


def test_brigade_level_in_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Бригада', 'Транспорт', 'Навыки', 'Уровень', 'Темп'])
    ws.append(['Б1', 'Авто', 'ремонт', 'Старший', None])
    ws.append(['Б2', 'Авто', 'ремонт', None, 0.85])
    p = tmp_path / 'b.xlsx'
    wb.save(p)
    assert [(b.level, b.speed) for b in read_brigades(str(p))] == [('старший', 0.9), ('', 0.85)]


def test_small_parsers():
    assert parse_minutes('17.08.2026 12:30') == 150
    assert parse_minutes('9:00') == -60
    assert parse_modes('') is None and parse_modes('любой') is None
    assert parse_modes('ОТ, пешком') == frozenset({'pt', 'foot'})
    assert parse_skills('L, c, Аварии') == [0, 1, 2]
    with pytest.raises(InputError):
        parse_skills('')


def test_xlsx_table(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Заявка', 'Начало', 'Число'])
    ws.append([7, datetime.datetime(2026, 8, 17, 14, 0), 3.0])
    ws.append([None, None, None])
    p = tmp_path / 't.xlsx'
    wb.save(p)
    assert read_table(str(p)) == [['Заявка', 'Начало', 'Число'], ['7', '17.08.2026 14:00', '3']]


def test_parse_window_bounds():
    from dispatch.files.readers import parse_window
    assert parse_window('17.08.2026 20:00', '17.08.2026 00:00') == (600, 840)        # «00:00» в конце — полночь
    assert parse_window('08:00', '11:00') == (-120, 60)                               # как в файле; смену учтёт задача
    for a, b in (('18:00', '16:00'), ('08:00', '09:00'), ('12:00', '12:00')):
        with pytest.raises(InputError):
            parse_window(a, b, 'строка 2')


def test_weekday_of_both_date_formats():
    from dispatch.config import weekday_of
    assert weekday_of('17.08.2026') == weekday_of('2026-08-17') == 0 and weekday_of('завтра') is None
