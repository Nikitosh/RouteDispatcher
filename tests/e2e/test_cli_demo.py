"""Сквозные тесты: настоящие решатели и кэши на демо-данных Востока, без OSRM (адреса серверов заведомо недоступны)."""
import json
import os
import subprocess
import sys

import openpyxl
import pytest

from dispatch.config import DATA_DIR, SOLVER_BIN

pytestmark = pytest.mark.e2e
ORDERS = os.path.join(DATA_DIR, 'demo', 'Восток заявки.xlsx')
BRIGADES = os.path.join(DATA_DIR, 'demo', 'Восток бригады.xlsx')
OFFLINE = {**os.environ, 'OSRM_CAR': 'http://127.0.0.1:9', 'OSRM_BIKE': 'http://127.0.0.1:9', 'OSRM_FOOT': 'http://127.0.0.1:9'}

if not all(os.path.exists(os.path.join(SOLVER_BIN, b)) for b in ('plan', 'baseline', 'fleet_bound')):
    pytest.skip('решатели не собраны: make -C solver', allow_module_level=True)


def cli(*args):
    p = subprocess.run([sys.executable, '-m', 'dispatch', *args], capture_output=True, text=True, env=OFFLINE, timeout=300)
    return p.returncode, p.stdout + p.stderr


@pytest.fixture(scope='module')
def planned(tmp_path_factory):
    out = tmp_path_factory.mktemp('vostok')
    code, text = cli('plan', ORDERS, '--name', 'Восток', '-b', BRIGADES, '-o', str(out))
    assert code == 0, text
    return out, json.load(open(out / 'plan.json')), text


def test_plan_is_valid_and_complete(planned):
    out, doc, text = planned
    m, base = doc['plan']['metrics'], doc['baseline']['metrics']
    assert doc['plan']['check'] == dict(ok=True, errors=[])
    assert m['served'] == m['orders'] == 66
    assert m['brigades_used'] == doc['lower_bound'] == 8            # минимум бригад доказан нижней границей
    assert m['brigades_used'] < base['brigades_used'] and m['km_total'] < base['km_total']
    assert all(o['explanation'] for o in doc['orders'] if o['plan'])
    # линии по улицам из кэша, без OSRM: решатель многопоточный, его план отличается от плана demo.sh, и части
    # переездов в кэше нет (такие маршруты — прямыми) — поэтому «хотя бы у половины»
    with_lines = sum(1 for r in doc['plan']['routes'] if r['geometry'])
    assert with_lines >= len(doc['plan']['routes']) / 2
    assert 'оптимален по числу бригад' in text


def test_durations_follow_complexity(planned):
    _, doc, text = planned
    O = doc['orders']
    assert all(o['detail'] for o in O) and len({o['svc'] for o in O if o['type'] == 'Подключение'}) > 3
    assert sum(1 for o in O if o['detail']['place']['src'].get('floor') == 'кадастр') > 40
    n = doc['norms']                                                # план по нормативу с реальной длительностью опаздывает
    assert n['work'] > n['work_norm'] and n['late'] > 0 and 'План по нормативу' in text


def test_car_only_orders_get_cars(planned):
    _, doc, _ = planned
    for o in doc['orders']:
        if o['need']:
            assert doc['brigades'][o['plan']['brigade']]['mode'] in o['need']


def test_xlsx_report(planned):
    out, _, _ = planned
    names = openpyxl.load_workbook(out / 'plan.xlsx').sheetnames
    assert names == ['Маршруты', 'Бригады', 'Метрики', 'Неназначенные', 'Объяснения', 'Базовый вариант', 'Базовый — неназначенные']


@pytest.mark.parametrize('event', [
    ['--unavailable', 'Бригада 4'],
    ['--cancel', '74198'],
    ['--urgent', '55.7155, 37.7625', '--type', 'подключение', '--window', '16:00-18:00'],
    ['--urgent', '55.7400, 37.7200', '--type', 'авария'],
])
def test_replan_events(planned, tmp_path, event):
    out, before, _ = planned
    code, text = cli('replan', str(out), '--at', '14:00', *event, '-o', str(tmp_path))
    assert code == 0, text
    doc = json.load(open(tmp_path / 'plan.json'))
    assert doc['plan']['check']['ok'], doc['plan']['check']['errors']
    at = doc['event']['at']
    old = {s['k']: s['beg'] for r in before['plan']['routes'] for s in r['stops']}
    new = {s['k']: s['beg'] for r in doc['plan']['routes'] for s in r['stops']}
    for k, start in old.items():                                     # начатое к 14:00 не сдвинулось
        if start < at and k in new:
            assert new[k] == start
    if '--type' in event and 'авария' in event:
        acc = doc['event']['accident']
        assert acc and acc['wait'] <= 120
    assert 'Изменения' in openpyxl.load_workbook(tmp_path / 'plan.xlsx').sheetnames


def test_replan_finish_later_and_strategies(planned, tmp_path):
    """В 14:00 работа, начатая раньше, затягивается на 2 часа: все три варианта ответа выполнимы."""
    out, before, _ = planned
    stop = max((s for r in before['plan']['routes'] for s in r['stops'] if s['beg'] < 240), key=lambda s: s['beg'])
    oid = before['orders'][stop['k']]['id']
    end = int(stop['end'] + 120)
    for strategy in ('insert', 'optimize', 'keep'):
        code, text = cli('replan', str(out), '--at', '14:00', '--finish', oid, f'{10 + end // 60}:{end % 60:02d}',
                         '--strategy', strategy, '--time', '1', '-o', str(tmp_path / strategy))
        assert code == 0, text
        doc = json.load(open(tmp_path / strategy / 'plan.json'))
        assert doc['plan']['check']['ok'], doc['plan']['check']['errors']
        assert doc['orders'][stop['k']]['plan']['end'] == end and doc['event']['delta'] == pytest.approx(120, abs=1)
        assert f'закончит заявку {oid}' in text


def test_plan_five_two_schedule_flat_traffic(tmp_path):
    code, text = cli('plan', ORDERS, '--name', 'Восток', '-b', BRIGADES, '--schedule', '5/2', '--flat-traffic',
                     '--time', '1', '--no-geometry', '-o', str(tmp_path))
    assert code == 0, text
    doc = json.load(open(tmp_path / 'plan.json'))
    shifts = [tuple(b['shift']) for b in doc['brigades']]
    assert set(shifts) == {(0, 540), (180, 720)} and shifts[:2] == [(0, 540), (180, 720)]
    assert doc['meta']['traffic_hours'] is False and 'Смены: 10:00–19:00' in text
    assert doc['plan']['check']['ok']
    for r in doc['plan']['routes']:                                   # каждая бригада — внутри своей смены
        b = doc['brigades'][r['v']]
        assert all(b['shift'][0] <= s['arr'] and s['end'] <= b['shift'][1] + 1e-6 for s in r['stops'])


def test_site(planned, tmp_path):
    out, _, _ = planned
    code, text = cli('site', str(out), '-o', str(tmp_path / 'site'))
    assert code == 0, text
    man = json.load(open(tmp_path / 'site' / 'manifest.json'))
    assert man['districts'][0]['name'] == 'Восток'


def test_bad_input_reports_line(tmp_path):
    bad = tmp_path / 'b.csv'
    bad.write_text('Бригада;Транспорт;Навыки\nБ1;самокат;ремонт\n', encoding='utf-8')
    code, text = cli('plan', ORDERS, '-b', str(bad), '-o', str(tmp_path / 'o'))
    assert code == 1 and 'строка 2: неизвестный транспорт «самокат»' in text
