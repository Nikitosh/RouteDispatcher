"""Примитивы слайдов в стиле «Билайна» (Beeline styles/шаблон презентации 2026.pptx).

Шаблон «Билайна» — 26,67 × 15 дюймов. Все координаты здесь задаются в сетке 13,33 × 7,5 (как у шаблона ЛЦТ)
и умножаются на SCALE, кегль — тоже. Так PDF обеих частей совпадает по пропорциям и размеру текста.
Шрифт — Beeline Sans (у начертания Black своё имя «Beeline Sans Black»). В шрифте нет «·» — не использовать.
"""
from lxml import etree
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SCALE = 2
BEE = RGBColor(0xFF, 0xC8, 0x00)       # Bee
BLACK = RGBColor(0x14, 0x14, 0x14)     # Line Black
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BEIGE = RGBColor(0xFF, 0xF9, 0xE5)     # фоновый Beige
LGREY = RGBColor(0xF1, 0xF1, 0xF1)     # фоновый Light Grey
GREY = RGBColor(0xA1, 0xA1, 0xA1)      # Grey
DGREY = RGBColor(0x5E, 0x5E, 0x5E)
MGREY = RGBColor(0x8C, 0x8C, 0x8C)
RED = RGBColor(0xE5, 0x09, 0x3B)       # дополнительный цвет темы
F_REG, F_BLACK = 'Beeline Sans', 'Beeline Sans Black'


def I(v):
    return Inches(v * SCALE)


def P(v):
    return Pt(v * SCALE)


def fill(sh, color):
    sh.fill.solid(); sh.fill.fore_color.rgb = color


def noline(sh):
    sh.line.fill.background()


def font(r, size, color, bold=False, black=False, italic=False):
    r.font.size = P(size); r.font.bold = bold and not black; r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = F_BLACK if black else F_REG


def set_bg(slide, color):
    cSld = slide._element.cSld
    for old in cSld.findall(qn('p:bg')):
        cSld.remove(old)
    xml = (f'<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
           f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:bgPr><a:solidFill>'
           f'<a:srgbClr val="{color}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>')
    cSld.insert(0, etree.fromstring(xml))


def _bullet(p, color, char='—'):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(I(0.24)))); pPr.set('indent', str(int(-I(0.24))))
    bc = etree.SubElement(pPr, qn('a:buClr')); c = etree.SubElement(bc, qn('a:srgbClr')); c.set('val', str(color))
    bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', F_REG)
    ch = etree.SubElement(pPr, qn('a:buChar')); ch.set('char', char)


def text(s, x, y, w, h, content, size=14, color=None, bold=False, black=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, space=6, line=None, bullet_color=None):
    """content: строка | список абзацев; абзац — строка или список кусков (текст, {опции})."""
    dark = getattr(s, 'kind', 'light') in ('dark', 'bee')
    if color is None:
        color = WHITE if getattr(s, 'kind', 'light') == 'dark' else BLACK
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i: p.space_before = P(space)
        if line: p.line_spacing = line
        chunks = para if isinstance(para, list) else [(para, {})]
        for chunk in chunks:
            t, o = (chunk, {}) if isinstance(chunk, str) else chunk
            r = p.add_run(); r.text = t
            font(r, o.get('size', size), o.get('color', color), o.get('bold', bold), o.get('black', black),
                 o.get('italic', False))
        if bullets:
            _bullet(p, bullet_color or (BEE if dark else BLACK))
    return tb


def box(s, x, y, w, h, color=LGREY, radius=0.05, outline=None):
    """Плашка: скруглённый прямоугольник без обводки, как в брендбуке."""
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = s.shapes.add_shape(shp, I(x), I(y), I(w), I(h))
    if radius:
        sh.adjustments[0] = radius
    fill(sh, color)
    if outline is not None:
        sh.line.color.rgb = outline; sh.line.width = P(0.75)
    else:
        noline(sh)
    sh.shadow.inherit = False
    sh.text_frame.text = ''
    return sh


def num(s, x, y, d, label, color=BEE, fg=BLACK, size=13):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(d), I(d))
    fill(sh, color); noline(sh); sh.shadow.inherit = False
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(label); font(r, size, fg, black=True)
    return sh


def arrow(s, x1, y1, x2, y2, color=BLACK, width=2.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = P(width)
    ln = c.line._get_or_add_ln()
    t = etree.SubElement(ln, qn('a:tailEnd')); t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')
    return c


def line(s, x1, y1, x2, y2, color=GREY, width=0.75, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = P(width)
    if dash:
        ln = c.line._get_or_add_ln()
        d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', 'dash')
    return c


def image(s, path_or_blob, x, y, w, h, size=None):
    """Картинка, вписанная в рамку с сохранением пропорций."""
    import io
    if isinstance(path_or_blob, bytes):
        src = io.BytesIO(path_or_blob); iw, ih = size or Image.open(io.BytesIO(path_or_blob)).size
    else:
        src = path_or_blob; iw, ih = Image.open(path_or_blob).size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    return s.shapes.add_picture(src, I(x + (w - pw) / 2), I(y + (h - ph) / 2), I(pw), I(ph))


def rounded(pic, radius=0.03):
    spPr = pic._element.spPr
    for g in spPr.findall(qn('a:prstGeom')):
        spPr.remove(g)
    g = etree.Element(qn('a:prstGeom')); g.set('prst', 'roundRect')
    av = etree.SubElement(g, qn('a:avLst')); gd = etree.SubElement(av, qn('a:gd'))
    gd.set('name', 'adj'); gd.set('fmla', f'val {int(radius * 100000)}')
    spPr.insert(1, g)
    return pic


def table(s, x, y, w, rows, col_w, row_h=0.4, size=11, align=None, bold_first_col=False, dark=False):
    shape = s.shapes.add_table(len(rows), len(rows[0]), I(x), I(y), I(w), I(row_h * len(rows)))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    for attr in ('bandRow', 'firstRow'):
        tblPr.set(attr, '0')
    st = tblPr.find(qn('a:tableStyleId'))
    if st is not None:
        st.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = I(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = I(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = I(0.1); c.margin_top = c.margin_bottom = I(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = BEE if dark else BLACK
            else:
                c.fill.fore_color.rgb = (RGBColor(0x26, 0x26, 0x26) if i % 2 == 0 else BLACK) if dark else \
                    (LGREY if i % 2 == 0 else WHITE)
            p = c.text_frame.paragraphs[0]; c.text_frame.word_wrap = True
            a = align[j] if align else ('l' if j == 0 else 'c')
            p.alignment = PP_ALIGN.LEFT if a == 'l' else PP_ALIGN.CENTER
            opts = {}
            if isinstance(val, tuple):
                val, opts = val
            r = p.add_run(); r.text = str(val)
            hdr_col = BLACK if dark else WHITE
            body_col = WHITE if dark else BLACK
            font(r, opts.get('size', size), opts.get('color', hdr_col if i == 0 else body_col),
                 bold=opts.get('bold', i == 0 or (bold_first_col and j == 0)))
            tcPr = c._tc.get_or_add_tcPr()
            for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                ln = etree.SubElement(tcPr, qn(tag)); ln.set('w', '0')
                etree.SubElement(ln, qn('a:noFill'))
            for f_ in [e for e in tcPr if e.tag in (qn('a:solidFill'), qn('a:noFill')) and e.getparent() is tcPr]:
                tcPr.remove(f_); tcPr.append(f_)
    return shape
