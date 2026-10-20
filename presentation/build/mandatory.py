"""Обязательные слайды 7–11 шаблона: правка текста и картинок без изменения вёрстки.

Правка на уровне run'ов: форматирование шаблона сохраняется.
"""
import copy
import io
import os

from PIL import Image
from pptx.util import Inches as I_, Pt

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
HERE = os.path.dirname(os.path.abspath(__file__))


def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def insert_pic(ph, image):
    """insert_picture берёт геометрию из макета — возвращаем геометрию плейсхолдера слайда."""
    geom = ph.left, ph.top, ph.width, ph.height
    pic = ph.insert_picture(image)
    pic.left, pic.top, pic.width, pic.height = geom
    return pic


def insert_round(ph, image):
    """Круглое фото: квадрат по высоте плейсхолдера, по центру его ширины."""
    left, top, width, height = ph.left, ph.top, ph.width, ph.height
    pic = ph.insert_picture(image)
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
    pic.left, pic.top, pic.width, pic.height = left + (width - height) // 2, top, height, height
    spPr = pic._element.spPr
    for g in spPr.findall(A + 'prstGeom'):
        spPr.remove(g)
    geom = spPr.makeelement(A + 'prstGeom', {'prst': 'ellipse'})
    geom.append(geom.makeelement(A + 'avLst', {}))
    spPr.insert(1, geom)
    return pic


def set_placeholder_text(sh, text):
    """Пустой плейсхолдер: один run, свойства берутся из endParaRPr."""
    p = sh.text_frame.paragraphs[0]
    end = p._p.find(A + 'endParaRPr')
    r = p.add_run()
    r.text = text
    if end is not None:
        rpr = copy.deepcopy(end)
        rpr.tag = A + 'rPr'
        old = r._r.find(A + 'rPr')
        if old is not None:
            r._r.remove(old)
        r._r.insert(0, rpr)


def fill_paragraphs(sh, items, proto_idx=0):
    """Заменяет абзацы текстового блока списком строк, копируя абзац-образец."""
    txb = sh.text_frame._txBody
    ps = sh.text_frame.paragraphs
    proto = copy.deepcopy(ps[proto_idx]._p)
    for p in ps:
        txb.remove(p._p)
    for text in items:
        p = copy.deepcopy(proto)
        runs = p.findall(A + 'r')
        for extra in runs[1:]:
            p.remove(extra)
        for br in p.findall(A + 'br'):
            p.remove(br)
        runs[0].find(A + 't').text = text
        txb.append(p)


def drop_paragraph(par):
    par._p.getparent().remove(par._p)


def set_size(shape, pt):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(pt)


import json as _json
_OUT = os.path.join(HERE, '..', '..', 'out')
_P = {n: _json.load(open(os.path.join(_OUT, n, 'plan.json'), encoding='utf-8'))
      for n in ('Восток', 'Югоцентр', 'Юго-восток', 'Север', 'Запад', 'Северо-восток')}
PROVEN = sum(1 for d in _P.values() if d.get('lower_bound') == d['plan']['metrics']['brigades_used'])
OURS3 = sum(_P[n]['plan']['metrics']['brigades_used'] for n in ('Восток', 'Югоцентр', 'Юго-восток'))
_pair = lambda n: (_P[n]['plan']['metrics']['brigades_used'], _P[n]['baseline']['metrics']['brigades_used'])
VOSTOK, YUG = _pair('Восток'), _pair('Югоцентр')
_m = lambda n, k: _P[n]['plan']['metrics'][k]
FULL = [n for n in _P if _m(n, 'served') == _m(n, 'orders')]
PART = [n for n in _P if n not in FULL]
_base = lambda n: _P[n]['baseline']['metrics']


def fill_mandatory(prs):
    S = prs.slides
    img = lambda name: os.path.join(HERE, name)

    # ---------- Слайд 7: титульный ----------
    s7 = S[6]
    set_placeholder_text(shape(s7, 740199916), 'Trdelnik')
    task = shape(s7, 757351073)
    set_placeholder_text(task, '3. Интеллектуальный сервис планирования рабочих маршрутов для инженеров '
                               'с учетом временных окон и сложности проводимых работ')
    set_size(task, 16)
    task.top = task.top - I_(0.4)
    task.text_frame.paragraphs[0]._p.get_or_add_pPr().set('indent', '0')
    task.text_frame.paragraphs[0]._p.get_or_add_pPr().set('marL', '268288')
    ph = shape(s7, 1636791008)
    pic = insert_pic(ph, img('beeline_business_white.png'))   # логотип с диска организаторов, перекрашен в белый
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
    w, h = Image.open(img('beeline_business_white.png')).size
    pic.width = int(pic.height * w / h)

    # ---------- Слайд 8: описание решения и команды ----------
    s8 = S[7]
    set_placeholder_text(shape(s8, 328740113), 'Команда Trdelnik')
    insert_pic(shape(s8, 1609484257), img('team_pair.jpg'))
    team = shape(s8, 661119785).text_frame.paragraphs
    team[0].runs[1].text = 'Никита Подгузов, бекенд-разработчик'
    team[1].runs[1].text = '2 человека'
    team[3].runs[0].text = 'вместе работаем в маленьком стартапе'
    drop_paragraph(team[4])
    team[5].runs[0].text = 'Город и регион: '
    r = team[5].add_run()
    r.text = 'Санкт-Петербург'
    r.font.bold = False
    r._r.get_or_add_rPr().append(copy.deepcopy(team[1].runs[1]._r.rPr.find(A + 'solidFill')))
    d = shape(s8, 1075761414)
    d.text_frame.paragraphs[0].runs[0].text = (
        'Сервис «Пульс» для диспетчера и инженера. Оптимизатор на C++ за секунды распределяет заявки участка: '
        'начальный план по упущенной выгоде, затем параллельные потоки (имитация отжига, ALNS, SISR) сокращают число '
        'бригад и пробег. Учитывает навыки, окна, смены 2/2 и 5/2, обед, оборудование, длительность работы по '
        'сложности дома и темпу бригады, старт из дома, рейтинг монтажников. Время в пути — дороги OSRM, пробки по '
        'часам и общественный транспорт по расписаниям. События дня пересчитываются от текущего положения бригад, '
        'в том числе разворот к аварии в пути. Экран инженера на телефоне.')
    set_size(d, 10.5)
    d.height = I_(1.85)                        # блок до разделительной линии: текст длиннее шаблонного
    bp = d.text_frame._txBody.find(A + 'bodyPr')
    for na in bp.findall(A + 'normAutofit'):
        bp.remove(na)
    u = shape(s8, 441518184)
    u.text_frame.paragraphs[0].runs[0].text = (
        f'Все заявки выполнены на {len(FULL)} участках из 6' + (', ' + ', '.join(
            f"{n} — {_m(n, 'served')} из {_m(n, 'orders')}" for n in PART) if PART else '') +
        '; меньше бригад нельзя — доказано ' + ('на всех шести' if PROVEN == 6 else f'на {PROVEN} из 6') +
        f'. На трёх участках организаторов {OURS3} бригад, в реальный день было 35. '
        f'Минимум бригад доказан в 88 из 88 тестовых задач в модели первого этапа, в 70 из 88 — в текущей.')

    # ---------- Слайд 9: состав команды (2 карточки из 5) ----------
    s9 = S[8]
    set_placeholder_text(shape(s9, 415042907), 'КОМАНДА TRDELNIK')
    for name_id, info_id, name, lines in [
            (1278750663, 1058462674, 'Никита Подгузов',
             ['Капитан, бекенд-разработчик', '@Nikitosh', '+7 993 478-72-39', 'Стартап Admin']),
            (1259743807, 990223453, 'Ксения Шор',
             ['Дизайнер', '@KseniiaShor', '+7 911 133-12-08', 'Стартап Admin'])]:
        shape(s9, name_id).text_frame.paragraphs[0].runs[0].text = name
        info = shape(s9, info_id).text_frame.paragraphs
        for p, t in zip(info, lines):
            p.runs[0].text = t
    photos = [insert_round(shape(s9, 1172799835), img('nikita_crop.jpg')),
              insert_round(shape(s9, 288949743), img('kseniia_crop.jpg'))]
    # две оставшиеся карточки ставим по центру слайда
    cards = [481545546, 1058462674, 1278750663, 1919445992, 990223453, 1259743807]
    left = min(shape(s9, i).left for i in cards)
    right = max(shape(s9, i).left + shape(s9, i).width for i in cards)
    dx = (prs.slide_width - (right - left)) // 2 - left
    for i in cards:
        shape(s9, i).left += dx
    for pic in photos:
        pic.left += dx
    for sid in [1717745144, 767261468, 789313918, 368022300,        # лишние карточки 3–5
                651037677, 499875170, 1054387886, 213007835,
                2086725175, 2106280066, 1456275104, 1344225912]:
        el = shape(s9, sid)._element
        el.getparent().remove(el)

    # ---------- Слайд 10: история команды ----------
    s10 = S[9]
    set_placeholder_text(shape(s10, 865297653), 'КОМАНДА TRDELNIK')
    shape(s10, 1588890124).text_frame.paragraphs[0].runs[0].text = (
        'Вместе работаем в маленьком стартапе, участвовали в ЛЦТ в 2024 году.')
    t = shape(s10, 1574440249).text_frame.paragraphs
    t[0].runs[0].text = ('Нам интересны задачи на оптимизацию. Здесь качество плана можно посчитать '
                         'и сравнить с тем, как день распределил диспетчер.')
    for p in t[1:]:
        drop_paragraph(p)
    fill_paragraphs(shape(s10, 243118211), [
        'Пришлось много работать с синтетическими данными и придумывать, какие сценарии по навыкам '
        'и транспорту нужно уметь поддерживать.',
    ])

    # ---------- Слайд 11: коротко о решении ----------
    s11 = S[10]
    fill_paragraphs(shape(s11, 1782181813), [
        'Время в пути: дороги OSRM, пробки по часам, дням недели и сетке ~330 м, общественный транспорт по '
        'расписаниям (автобусы, метро, МЦК, МЦД, электрички).',
        'Длительность работы по сложности: этаж, лифт, год постройки дома, тип работы и темп бригады; коэффициенты '
        'уточняются по историческим данным о выездах.',
        'Оптимизатор на C++: параллельные потоки (отжиг, ALNS, SISR) с обедом, сменами, оборудованием, рейтингом; '
        'обычный участок за 3 секунды. Нижняя граница — генерация столбцов.',
        'Пересчёт событий дня от текущего положения бригад, независимая проверка каждого плана.',
    ])
    fill_paragraphs(shape(s11, 321655550), [
        f"Меньше бригад и больше выполненных заявок: на Востоке {VOSTOK[0]} бригад выполняют {_m('Восток', 'served')} из "
        f"{_m('Восток', 'orders')} заявок, базовый вариант — {VOSTOK[1]} бригад и {_base('Восток')['served']}; "
        f"на Югоцентре {YUG[0]} и {_m('Югоцентр', 'served')} против {YUG[1]} и {_base('Югоцентр')['served']}.",
        'Диспетчер видит маршруты, причины назначений и последствия любого события; инженер — свой день на телефоне.',
        'Руководитель видит, сколько бригад нужно, во что обходятся правила и где не хватает людей.',
        'На вход нужна только выгрузка заявок и бригад в CSV или xlsx.',
        'Дальше: планирование на несколько дней, пополнение оборудования днём.',
    ])
