"""Примитивы для свободных слайдов в стиле шаблона: фон, плашка заголовка, карточки, текст, стрелки."""
import copy

from lxml import etree
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

PINK = RGBColor(0xFF, 0x00, 0x53)
PURPLE = RGBColor(0x52, 0x09, 0x78)
DARK = RGBColor(0x31, 0x0F, 0x53)
LAV = RGBColor(0x8A, 0x83, 0xD1)
ROSE = RGBColor(0xFF, 0xD6, 0xE4)
INK = RGBColor(0x1C, 0x1D, 0x22)
MUTED = RGBColor(0x6B, 0x68, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xE9, 0xE2, 0xF6)      # второстепенный текст на тёмном фоне
TINT = RGBColor(0xF6, 0xF3, 0xFB)      # фон карточки на светлом слайде

I = Inches


class Deck:
    def __init__(self, prs, layout, bgs):
        self.prs, self.layout, self.bgs = prs, layout, bgs
        self.new = []

    def slide(self, kind, label, headline=None):
        """kind: 'dark' | 'light'. label — короткая плашка, headline — вывод слайда."""
        s = self.prs.slides.add_slide(self.layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        set_bg(s, self.bgs[kind])
        s.kind = kind
        w = 0.62 + 0.205 * len(label)
        pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(0.38), I(0.35), I(w), I(0.68))
        pill.adjustments[0] = 0.187
        fill(pill, PINK); noline(pill)
        tf = pill.text_frame
        tf.margin_left = tf.margin_right = I(0.2); tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = label.upper(); font(r, 20, WHITE, bold=True)
        if headline:
            text(s, 0.4, 1.22, 12.5, 0.62, headline, size=24, bold=True,
                 color=WHITE if kind == 'dark' else PURPLE)
        self.new.append(s)
        return s


def set_bg(slide, img_part):
    rId = slide.part.relate_to(img_part, RT.IMAGE)
    xml = (f'<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
           f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
           f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
           f'<p:bgPr><a:blipFill rotWithShape="1"><a:blip r:embed="{rId}"/><a:stretch><a:fillRect/></a:stretch>'
           f'</a:blipFill><a:effectLst/></p:bgPr></p:bg>')
    slide._element.cSld.insert(0, etree.fromstring(xml))


def bg_part(slide):
    blip = slide._element.cSld.find(qn('p:bg')).find('.//' + qn('a:blip'))
    return slide.part.related_part(blip.get(qn('r:embed')))


def fill(sh, color):
    sh.fill.solid(); sh.fill.fore_color.rgb = color


def noline(sh):
    sh.line.fill.background()


def font(r, size, color, bold=False, italic=False):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color


def _bullet(p, color):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(I(0.2)))); pPr.set('indent', str(int(-I(0.2))))
    for tag in ('a:buClr', 'a:buFont', 'a:buChar', 'a:buNone'):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    bc = etree.SubElement(pPr, qn('a:buClr')); c = etree.SubElement(bc, qn('a:srgbClr')); c.set('val', str(color))
    bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', 'Arial')
    ch = etree.SubElement(pPr, qn('a:buChar')); ch.set('char', '•')


def text(s, x, y, w, h, content, size=14, color=None, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, space=6, line=None):
    """content: строка | список абзацев; абзац — строка или список кусков (текст, {опции})."""
    if color is None:
        color = WHITE if getattr(s, 'kind', 'light') == 'dark' else INK
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i: p.space_before = Pt(space)
        if line: p.line_spacing = line
        chunks = para if isinstance(para, list) else [(para, {})]
        for chunk in chunks:
            t, o = (chunk, {}) if isinstance(chunk, str) else chunk
            r = p.add_run(); r.text = t
            font(r, o.get('size', size), o.get('color', color), o.get('bold', bold), o.get('italic', False))
        if bullets:
            _bullet(p, PINK)
    return tb


def card(s, x, y, w, h, color=None, outline=None, radius=0.06):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    sh.adjustments[0] = radius
    if color is None:
        color = WHITE
    fill(sh, color)
    if outline is None and getattr(s, 'kind', 'light') == 'light' and color == WHITE:
        outline = LAV
    if outline is not None:
        sh.line.color.rgb = outline; sh.line.width = Pt(0.9)
    else:
        noline(sh)
    sh.shadow.inherit = False
    sh.text_frame.text = ''
    return sh


def circle_num(s, x, y, d, label, color=PINK, fg=WHITE, size=14):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(d), I(d))
    fill(sh, color); noline(sh); sh.shadow.inherit = False
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(label); font(r, size, fg, bold=True)
    return sh


def pill_tag(s, x, y, label, color=ROSE, fg=PURPLE, size=10.5, h=0.3):
    w = 0.24 + 0.085 * len(label) * size / 10.5
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(y), I(w), I(h))
    sh.adjustments[0] = 0.5; fill(sh, color); noline(sh); sh.shadow.inherit = False
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; font(r, size, fg, bold=True)
    return w


def arrow(s, x1, y1, x2, y2, color=PINK, width=2.0, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', 'dash')
    t = etree.SubElement(ln, qn('a:tailEnd')); t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')
    return c


def line(s, x1, y1, x2, y2, color=LAV, width=1.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    return c


def image(s, path, x, y, w, h, align='center'):
    """Картинка, вписанная в рамку с сохранением пропорций."""
    iw, ih = Image.open(path).size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    px = x + (w - pw) / 2 if align == 'center' else x
    py = y + (h - ph) / 2
    return s.shapes.add_picture(path, I(px), I(py), I(pw), I(ph))


def rounded_image(s, path, x, y, w, h, radius=0.03):
    pic = image(s, path, x, y, w, h)
    spPr = pic._element.spPr
    for g in spPr.findall(qn('a:prstGeom')):
        spPr.remove(g)
    g = etree.SubElement(spPr, qn('a:prstGeom')); g.set('prst', 'roundRect')
    av = etree.SubElement(g, qn('a:avLst')); gd = etree.SubElement(av, qn('a:gd'))
    gd.set('name', 'adj'); gd.set('fmla', f'val {int(radius * 100000)}')
    spPr.remove(g); spPr.insert(1, g)
    return pic


def big_stat(s, x, y, w, value, label, vcolor=None, lcolor=None, vsize=40, lsize=12):
    dark = getattr(s, 'kind', 'light') == 'dark'
    text(s, x, y, w, 0.75, value, size=vsize, bold=True, color=vcolor or (WHITE if dark else PINK))
    text(s, x, y + 0.08 + vsize / 60, w, 0.7, label, size=lsize, color=lcolor or (SOFT if dark else MUTED))


def page_numbers(prs):
    for i, s in enumerate(prs.slides, 1):
        if getattr(s, 'kind', None) is None:
            continue
        text(s, 12.2, 6.98, 0.85, 0.3, str(i), size=10, align=PP_ALIGN.RIGHT,
             color=WHITE if s.kind == 'dark' else PURPLE)


def table(s, x, y, w, rows, col_w, row_h=0.36, size=11, header_fill=PURPLE, zebra=TINT, bold_first_col=False, align=None):
    shape = s.shapes.add_table(len(rows), len(rows[0]), I(x), I(y), I(w), I(row_h * len(rows)))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    for attr in ('bandRow', 'firstRow'):
        tblPr.set(attr, '0')
    style = tblPr.find(qn('a:tableStyleId'))
    if style is not None:
        style.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'   # «Без стиля, сетка таблицы»
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = I(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = I(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = I(0.08); c.margin_top = c.margin_bottom = I(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = header_fill if i == 0 else (zebra if i % 2 == 0 else WHITE)
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            a = align[j] if align else ('l' if j == 0 else 'c')
            p.alignment = PP_ALIGN.LEFT if a == 'l' else PP_ALIGN.CENTER
            opts = {}
            if isinstance(val, tuple):
                val, opts = val
            r = p.add_run(); r.text = str(val)
            font(r, opts.get('size', size), opts.get('color', WHITE if i == 0 else INK),
                 bold=opts.get('bold', i == 0 or (bold_first_col and j == 0)))
            # тонкие светлые границы
            tcPr = c._tc.get_or_add_tcPr()
            for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                ln = etree.SubElement(tcPr, qn(tag)); ln.set('w', '6350')
                sf = etree.SubElement(ln, qn('a:solidFill')); cl = etree.SubElement(sf, qn('a:srgbClr'))
                cl.set('val', 'E2DFEC')
            # элементы tcPr по схеме: границы идут до заливки
            fills = [e for e in tcPr if e.tag in (qn('a:solidFill'), qn('a:noFill'))]
            for f_ in fills:
                tcPr.remove(f_); tcPr.append(f_)
    return shape


def drop_slides(prs, keep):
    """Оставить только слайды с индексами keep (в этом порядке); остальные удалить вместе со связями."""
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    keep_ids = [ids[i] for i in keep]
    for sid in ids:
        if sid not in keep_ids:
            prs.part.drop_rel(sid.get(qn('r:id')))
            sldIdLst.remove(sid)
    for sid in keep_ids:
        sldIdLst.remove(sid)
    for sid in keep_ids:
        sldIdLst.append(sid)
