"""plan.xlsx для диспетчера: маршруты, бригады, метрики, неназначенные, объяснения, базовый вариант, изменения."""
import os

from ..formatting import hhmm

ROUTE_HEADER = ['Бригада', 'Транспорт', '№', 'Заявка', 'Тип', 'Район', 'Адрес', 'Окно', 'Приезд', 'Начало', 'Конец', 'Км от предыдущей']
ROUTE_WIDTHS = [14, 16, 5, 12, 13, 16, 44, 13, 9, 9, 9, 10]


def _window(o):
    return f"{hhmm(o['window'][0])}–{hhmm(o['window'][1])}"


class XlsxReport:
    def __init__(self, doc):
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        self.doc = doc
        self.wb = openpyxl.Workbook()
        self.bold, self.fill = Font(bold=True), PatternFill('solid', fgColor='EDEBE6')

    def sheet(self, title, header, rows, widths):
        from openpyxl.styles import Alignment
        from openpyxl.utils import get_column_letter
        ws = self.wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = self.bold, self.fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical='top', wrap_text=isinstance(c.value, str) and len(c.value) > 40)
        ws.freeze_panes = 'A2'
        return ws

    def route_rows(self, routes):
        B, O = self.doc['brigades'], self.doc['orders']
        return [[B[r['v']]['name'], B[r['v']]['mode_name'], i, O[s['k']]['id'], O[s['k']]['type'], O[s['k']]['district'],
                 O[s['k']]['addr'], _window(O[s['k']]), hhmm(s['arr']), hhmm(s['beg']), hhmm(s['end']), round(s['km'], 2)]
                for r in routes for i, s in enumerate(r['stops'], 1)]

    def build(self):
        d, B, O = self.doc, self.doc['brigades'], self.doc['orders']
        if d.get('event'):
            self.changes_sheet()
        self.sheet('Маршруты', ROUTE_HEADER, self.route_rows(d['plan']['routes']), ROUTE_WIDTHS)
        used = {r['v']: r for r in d['plan']['routes']}
        self.sheet('Бригады', ['Бригада', 'Транспорт', 'Навыки', 'Старт', 'Заявок', 'Первая', 'Последняя до', 'Км'],
                   [[b['name'], b['mode_name'], ', '.join(b['skills']), b['start_addr'],
                     len(used[b['v']]['stops']) if b['v'] in used else 0,
                     hhmm(used[b['v']]['stops'][0]['beg']) if b['v'] in used else '—',
                     hhmm(used[b['v']]['stops'][-1]['end']) if b['v'] in used else '—',
                     used[b['v']]['km'] if b['v'] in used else 0] for b in B], [14, 16, 34, 44, 8, 9, 13, 8])
        self.metrics_sheet()
        self.sheet('Неназначенные', ['Заявка', 'Тип', 'Район', 'Адрес', 'Окно', 'Причина'],
                   [[o['id'], o['type'], o['district'], o['addr'], _window(o), o['reason']['text']]
                    for o in O if o['plan'] is None and o.get('reason')] or [['Все заявки назначены']], [12, 13, 16, 44, 13, 70])
        self.sheet('Объяснения', ['Заявка', 'Тип', 'Район', 'Бригада', 'Почему эта бригада', 'Почему не другие'],
                   [[o['id'], o['type'], o['district'], B[o['plan']['brigade']]['name'], o['explanation']['summary'],
                     '\n'.join(f"{g['names']} — {g['text']}" for g in o['explanation']['others'])]
                    for o in O if o['plan'] and o.get('explanation')], [12, 13, 16, 14, 60, 80])
        self.sheet('Базовый вариант', ROUTE_HEADER, self.route_rows(d['baseline']['routes']), ROUTE_WIDTHS)
        self.sheet('Базовый — неназначенные', ['Заявка', 'Тип', 'Район', 'Адрес', 'Окно', 'Причина'],
                   [[O[u['k']]['id'], O[u['k']]['type'], O[u['k']]['district'], O[u['k']]['addr'], _window(O[u['k']]), u['text']]
                    for u in d['baseline']['unassigned']] or [['Все заявки назначены']], [12, 13, 16, 44, 13, 70])
        del self.wb['Sheet']
        return self.wb

    def metrics_sheet(self):
        m, mb, d = self.doc['plan']['metrics'], self.doc['baseline']['metrics'], self.doc
        delay = lambda x: '—' if x is None else f'{x:.0f} мин'
        rows = [['Задействовано бригад', m['brigades_used'], mb['brigades_used']],
                ['Пробег всего, км', m['km_total'], mb['km_total']],
                ['Выполнено заявок', f"{m['served']} из {m['orders']}", f"{mb['served']} из {mb['orders']}"],
                ['Неназначенных', m['unassigned'], mb['unassigned']],
                ['Авария: среднее ожидание от начала окна', delay(m['accident_delay_avg']), delay(mb['accident_delay_avg'])],
                ['Авария: максимальное ожидание', delay(m['accident_delay_max']), delay(mb['accident_delay_max'])],
                ['Нижняя граница числа бригад', d['lower_bound'] if d['lower_bound'] is not None else '—', ''],
                ['Проверка ограничений', 'нарушений нет' if d['plan']['check']['ok'] else '; '.join(d['plan']['check']['errors']), ''],
                [], ['Пробег по бригадам, км', 'Наш план', 'Базовый вариант']]
        names = list(dict.fromkeys(list(m['km_by_brigade']) + list(mb['km_by_brigade'])))
        rows += [[n, m['km_by_brigade'].get(n, '—'), mb['km_by_brigade'].get(n, '—')] for n in names]
        ws = self.sheet('Метрики', ['Показатель', 'Наш план', 'Базовый вариант ТЗ'], rows, [42, 18, 20])
        ws.cell(row=len(rows) - len(names) + 1, column=1).font = self.bold

    def changes_sheet(self):
        e, B, O = self.doc['event'], self.doc['brigades'], self.doc['orders']
        slot = lambda x: f"{B[x['brigade']]['name']} {hhmm(x['start'])}" if x else '—'
        rows = [[O[c['k']]['id'], O[c['k']]['type'], O[c['k']]['district'], slot(c['was']),
                 slot(c['now']) if c['now'] else 'не назначена', c['why']] for c in e['changes']]
        ws = self.sheet('Изменения', ['Заявка', 'Тип', 'Район', 'Было', 'Стало', 'Причина'], rows or [['Изменений нет']],
                        [12, 13, 16, 20, 20, 50])
        ws.insert_rows(1, 2)
        ws['A1'] = f"Событие в {hhmm(e['at'])}; зафиксировано выполненных и начатых заявок: {e['fixed']}"
        ws['A1'].font = self.bold


def write_xlsx(doc, out_dir):
    path = os.path.join(out_dir, 'plan.xlsx')
    XlsxReport(doc).build().save(path)
    return path
