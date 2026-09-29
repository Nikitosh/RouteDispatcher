"""Файлы бригад и заявок для демо по данным организаторов (data/demo/).

Бригады — состав контрольного дня: навыки — типы заявок, которые бригада реально выполняла; бригадам, работавшим в
основном в Кашире, Ступино или Домодедово, старт в их городе. Транспорт задаём сами (жюри разрешило): все на общ.
транспорте, все на машине или смесь (mix1, mix2 — воспроизводимо по имени бригады).
Заявки — синтетика организаторов плюс колонка «Транспорт»: только машина для настоящей аварии на узле связи
(«Глобальная проблема / Авария»: лестница, сварка оптики) и для подключения с заказом оборудования
(«Подключение / Заказ подключения…»: роутер, приставка, кабель). Остальным — любой транспорт.
"""
import os
import random
import zlib
from collections import Counter

from ..config import DATA_DIR, MODE_NAMES, SKILL, SKILL_NAMES
from ..files.readers import read_orders
from ..files.tables import read_table
from ..geo.geocoder import Geocoder
from .sources import organizers_file

FAR_TOWNS = ('Кашира', 'Ступино', 'Домодедово')
REGIONS = ('Восток', 'Юго-восток', 'Югоцентр')
DEMO_DIR = os.path.join(DATA_DIR, 'demo')


def mode_of(name, far, variant):
    if variant in ('pt', 'car'):
        return variant
    if far:
        return 'car'
    r = random.Random(zlib.crc32(f'{name}|{variant[-1]}'.encode())).random()
    return 'pt' if r < 0.60 else 'car' if r < 0.85 else 'bike' if r < 0.95 else 'foot'


def needs_car(bk, hd):
    return (bk == 'Глобальная проблема' and hd == 'Авария') or (bk == 'Подключение' and hd.startswith('Заказ подключения'))


def roster(region, variant='mix2', geocoder=None):
    """[(имя, транспорт, навыки, старт)] по контрольному дню участка; порядок — по числу заявок бригады."""
    orders = read_orders(organizers_file(region, 'Синтетические')).orders
    rows = read_table(organizers_file(region, 'Контрольное'))
    col = rows[0].index('Бригада')
    day = [r for r in rows[1:] if r[0]]
    coords = (geocoder or Geocoder()).locate([o.address for o in orders], log=lambda *_: None)
    out = []
    for n, (name, _) in enumerate(Counter(r[col] for r in day if r[col]).most_common(), 1):
        mine = [o for r, o in zip(day, orders) if r[col] == name]
        town, cnt = Counter(o.district for o in mine).most_common(1)[0]
        start = ''
        if town in FAR_TOWNS and cnt * 2 >= len(mine):
            pts = [coords[o.address] for o in mine if o.district == town]
            start = f'{sum(p[0] for p in pts) / len(pts):.6f}, {sum(p[1] for p in pts) / len(pts):.6f}'
        skills = sorted({SKILL[o.bk] for o in mine})
        out.append((f'Бригада {n}', mode_of(name, bool(start), variant), skills, start))
    return out


def write_brigades(path, rows):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Бригады'
    ws.append(['Бригада', 'Транспорт', 'Навыки', 'Старт'])
    for c in ws[1]:
        c.font = Font(bold=True)
    for name, mode, skills, start in rows:
        ws.append([name, MODE_NAMES[mode], ', '.join(SKILL_NAMES[s] for s in skills), start])
    for col, w in zip('ABCD', (14, 18, 42, 24)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def write_orders(synth_csv, path):
    """Синтетика организаторов → xlsx с колонкой «Транспорт». Возвращает (заявок, из них «только машина»)."""
    import openpyxl
    from openpyxl.styles import Font
    rows = read_table(synth_csv)
    head = rows[0]
    bk, hd = head.index('Тип заявки BK'), head.index('Тип заявки HD')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Заявки'
    ws.append(head + ['Транспорт'])
    for c in ws[1]:
        c.font = Font(bold=True)
    total = need = 0
    for r in rows[1:]:
        if r[0].lower().startswith('адрес'):
            ws.append(r)
            continue
        car = needs_car(r[bk], r[hd])
        ws.append(r + [''] * (len(head) - len(r)) + ['авто' if car else ''])
        total += 1
        need += car
    for col, w in zip('ABCDEFGHIJ', (11, 20, 30, 17, 17, 15, 44, 12, 12, 11)):
        ws.column_dimensions[col].width = w
    wb.save(path)
    return total, need


def write_all(variant='mix2', out_dir=DEMO_DIR):
    os.makedirs(out_dir, exist_ok=True)
    done = []
    for region in REGIONS:
        b = os.path.join(out_dir, f'{region} бригады.xlsx')
        o = os.path.join(out_dir, f'{region} заявки.xlsx')
        rows = roster(region, variant)
        write_brigades(b, rows)
        total, need = write_orders(organizers_file(region, 'Синтетические'), o)
        done.append((region, b, o, len(rows), total, need))
    return done
